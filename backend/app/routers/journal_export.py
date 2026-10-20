"""
Trade journal export — Excel (.xlsx) and PDF.

GET /api/journal/export/excel
GET /api/journal/export/pdf
    optional query params: start=YYYY-MM-DD, end=YYYY-MM-DD, pair=GBP/USD, result=WIN|LOSS|BREAKEVEN|OPEN

Excel workbook:
  * Summary  — headline stats as live formulas over the Trades sheet, plus a
               per-pair and per-timeframe breakdown (also formulas)
  * Trades   — every trade, one row each, as a filterable Excel table
  * Equity   — running P&L after each closed trade, with a line chart
PDF report: summary, per-pair table, equity curve and the full trade list.

Needs `openpyxl` and `reportlab` in requirements.txt.
"""
import io
import re
from datetime import datetime, date, time, timezone
from typing import Optional
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.database import get_db
from app.middleware.auth_middleware import get_current_user
from app.models.trade import Trade
from app.models.user import User

router = APIRouter(prefix="/api/journal/export", tags=["Trade Journal"])

CLOSED = ("WIN", "LOSS", "BREAKEVEN")


# --------------------------------------------------------------------- data
def _tz(user: User):
    try:
        return ZoneInfo(user.timezone or "Africa/Dar_es_Salaam")
    except Exception:
        return ZoneInfo("Africa/Dar_es_Salaam")


def _local(dt, tz):
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(tz).replace(tzinfo=None)  # Excel has no timezones


def _row(t: Trade, tz) -> dict:
    notes = t.user_notes or ""
    auto = notes.startswith("[AUTO]")
    tf = re.search(r"\[AUTO\]\s*([A-Z0-9?]+)", notes)
    reason = re.search(r"closed by ([A-Z ]+?)(?:\s·|$)", notes)
    risk = float(t.risk_amount or 0)
    pnl = float(t.pnl) if t.pnl is not None else None
    r_mult = round(pnl / risk, 2) if (pnl is not None and risk > 0) else None
    return {
        "opened": _local(t.opened_at, tz),
        "closed": _local(t.closed_at, tz),
        "pair": t.pair,
        "direction": t.direction,
        "timeframe": tf.group(1) if tf else "",
        "source": "Auto (MT5)" if auto else "Manual",
        "entry": t.entry_price,
        "stop_loss": t.stop_loss,
        "take_profit": t.take_profit,
        "close": t.close_price,
        "lots": t.lot_size,
        "risk": round(risk, 2),
        "result": t.result,
        "pnl": pnl,
        "pips": t.pnl_pips,
        "r": r_mult,
        "close_reason": reason.group(1).title() if reason else "",
        "notes": notes,
    }


def load_rows(db: Session, user: User, start: Optional[date], end: Optional[date],
              pair: Optional[str], result: Optional[str]) -> list:
    tz = _tz(user)
    q = db.query(Trade).filter(Trade.user_id == user.id)
    if pair:
        q = q.filter(Trade.pair == pair.strip().upper().replace("-", "/"))
    if result:
        q = q.filter(Trade.result == result.strip().upper())
    if start:
        q = q.filter(Trade.opened_at >= datetime.combine(start, time.min, tz))
    if end:
        q = q.filter(Trade.opened_at <= datetime.combine(end, time.max, tz))
    return [_row(t, tz) for t in q.order_by(Trade.opened_at.asc()).all()]


def summarize(rows: list) -> dict:
    closed = [r for r in rows if r["result"] in CLOSED and r["pnl"] is not None]
    wins = [r for r in closed if r["result"] == "WIN"]
    losses = [r for r in closed if r["result"] == "LOSS"]
    gross_win = sum(r["pnl"] for r in wins)
    gross_loss = -sum(r["pnl"] for r in losses)
    equity, peak, max_dd = 0.0, 0.0, 0.0
    streak = worst_streak = 0
    for r in sorted(closed, key=lambda x: x["closed"] or x["opened"] or datetime.min):
        equity += r["pnl"]
        peak = max(peak, equity)
        max_dd = max(max_dd, peak - equity)
        streak = streak + 1 if r["result"] == "LOSS" else 0
        worst_streak = max(worst_streak, streak)
    rs = [r["r"] for r in closed if r["r"] is not None]
    return {
        "total": len(rows),
        "open": sum(1 for r in rows if r["result"] == "OPEN"),
        "closed": len(closed),
        "wins": len(wins),
        "losses": len(losses),
        "win_rate": (len(wins) / len(closed) * 100) if closed else 0.0,
        "net": sum(r["pnl"] for r in closed),
        "avg_win": (gross_win / len(wins)) if wins else 0.0,
        "avg_loss": (-gross_loss / len(losses)) if losses else 0.0,
        "profit_factor": (gross_win / gross_loss) if gross_loss > 0 else None,
        "avg_r": (sum(rs) / len(rs)) if rs else None,
        "max_dd": max_dd,
        "worst_streak": worst_streak,
    }


def _period(rows, start, end) -> str:
    if start or end:
        return f"{start or 'start'} to {end or 'today'}"
    if rows:
        first = rows[0]["opened"]
        return f"{first:%d %b %Y} to {datetime.now():%d %b %Y}" if first else "All trades"
    return "All trades"


def _filename(ext: str) -> str:
    return f"forex_intel_journal_{datetime.now():%Y-%m-%d}.{ext}"


# -------------------------------------------------------------------- Excel
TRADE_COLS = [
    ("Opened", "opened", 18, "yyyy-mm-dd hh:mm"),
    ("Closed", "closed", 18, "yyyy-mm-dd hh:mm"),
    ("Pair", "pair", 10, None),
    ("Direction", "direction", 10, None),
    ("Timeframe", "timeframe", 10, None),
    ("Source", "source", 12, None),
    ("Entry", "entry", 11, "0.00000"),
    ("Stop loss", "stop_loss", 11, "0.00000"),
    ("Take profit", "take_profit", 11, "0.00000"),
    ("Close price", "close", 11, "0.00000"),
    ("Lots", "lots", 7, "0.00"),
    ("Risk ($)", "risk", 9, "#,##0.00"),
    ("Result", "result", 11, None),
    ("P&L ($)", "pnl", 10, "#,##0.00;[Red]-#,##0.00"),
    ("Pips", "pips", 8, "0.0;[Red]-0.0"),
    ("R multiple", "r", 10, "0.00;[Red]-0.00"),
    ("Closed by", "close_reason", 15, None),
    ("Notes", "notes", 60, None),
]


def build_excel(rows: list, period: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF", name="Arial")
    head_fill = PatternFill("solid", start_color="1F2A44")
    base_font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)

    # ---- Trades sheet (data first; Summary formulas point at it)
    ws_t = wb.active
    ws_t.title = "Trades"
    col_of = {}
    for i, (title, key, width, _) in enumerate(TRADE_COLS, start=1):
        c = ws_t.cell(row=1, column=i, value=title)
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(horizontal="center")
        ws_t.column_dimensions[get_column_letter(i)].width = width
        col_of[key] = get_column_letter(i)
    for r_i, row in enumerate(rows, start=2):
        for c_i, (_, key, _, fmt) in enumerate(TRADE_COLS, start=1):
            c = ws_t.cell(row=r_i, column=c_i, value=row[key])
            c.font = base_font
            if fmt:
                c.number_format = fmt
    last = max(len(rows) + 1, 2)
    if rows:
        tbl = Table(displayName="Trades", ref=f"A1:{get_column_letter(len(TRADE_COLS))}{last}")
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws_t.add_table(tbl)
    ws_t.freeze_panes = "A2"

    res = f"Trades!${col_of['result']}$2:${col_of['result']}${last}"
    pnl = f"Trades!${col_of['pnl']}$2:${col_of['pnl']}${last}"
    rcol = f"Trades!${col_of['r']}$2:${col_of['r']}${last}"
    pair_rng = f"Trades!${col_of['pair']}$2:${col_of['pair']}${last}"
    tf_rng = f"Trades!${col_of['timeframe']}$2:${col_of['timeframe']}${last}"

    # ---- Summary sheet (live formulas, so it updates if the user edits Trades)
    ws = wb.create_sheet("Summary", 0)
    ws.column_dimensions["A"].width = 26
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 14
    ws["A1"] = "Forex Intel — Trade Journal"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = f"Period: {period}   ·   Exported {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="666666")

    stats = [
        ("Total trades", f"=COUNTA({res})", "0"),
        ("Open trades", f'=COUNTIF({res},"OPEN")', "0"),
        ("Closed trades", '=B7+B8+B9', "0"),
        ("Wins", f'=COUNTIF({res},"WIN")', "0"),
        ("Losses", f'=COUNTIF({res},"LOSS")', "0"),
        ("Breakeven", f'=COUNTIF({res},"BREAKEVEN")', "0"),
        ("Win rate", '=IF(B6>0,B7/B6,0)', "0.0%"),
        ("Net P&L ($)", f'=SUMIF({res},"<>OPEN",{pnl})', "#,##0.00;[Red]-#,##0.00"),
        ("Average win ($)", f'=IFERROR(AVERAGEIF({res},"WIN",{pnl}),0)', "#,##0.00"),
        ("Average loss ($)", f'=IFERROR(AVERAGEIF({res},"LOSS",{pnl}),0)', "#,##0.00;[Red]-#,##0.00"),
        ("Profit factor", f'=IFERROR(SUMIF({res},"WIN",{pnl})/-SUMIF({res},"LOSS",{pnl}),0)', "0.00"),
        ("Average R per trade", f'=IFERROR(AVERAGE({rcol}),0)', "0.00;[Red]-0.00"),
        ("Max drawdown ($)", "=IFERROR(MAX(Equity!F:F),0)", "#,##0.00"),
    ]
    for i, (label, formula, fmt) in enumerate(stats, start=4):
        ws.cell(row=i, column=1, value=label).font = base_font
        c = ws.cell(row=i, column=2, value=formula)
        c.font, c.number_format = bold, fmt
    ws["D4"] = "Break-even win rate at 2:1 is about 33%, a bit higher after costs."
    ws["D4"].font = Font(name="Arial", size=9, italic=True, color="666666")

    def breakdown(start_row: int, title: str, rng: str, values: list) -> int:
        hdr = [title, "Trades", "Wins", "Losses", "Win rate", "Net P&L ($)", "Avg R"]
        for j, h in enumerate(hdr, start=1):
            c = ws.cell(row=start_row, column=j, value=h)
            c.font, c.fill = head_font, head_fill
        for k, v in enumerate(values, start=start_row + 1):
            ws.cell(row=k, column=1, value=v).font = base_font
            ws.cell(row=k, column=2, value=f'=COUNTIFS({rng},A{k},{res},"<>OPEN")')
            ws.cell(row=k, column=3, value=f'=COUNTIFS({rng},A{k},{res},"WIN")')
            ws.cell(row=k, column=4, value=f'=COUNTIFS({rng},A{k},{res},"LOSS")')
            c = ws.cell(row=k, column=5, value=f"=IF(B{k}>0,C{k}/B{k},0)")
            c.number_format = "0.0%"
            c = ws.cell(row=k, column=6, value=f'=SUMIFS({pnl},{rng},A{k},{res},"<>OPEN")')
            c.number_format = "#,##0.00;[Red]-#,##0.00"
            c = ws.cell(row=k, column=7, value=f'=IFERROR(AVERAGEIFS({rcol},{rng},A{k},{res},"<>OPEN"),0)')
            c.number_format = "0.00;[Red]-0.00"
            for col in range(2, 8):
                ws.cell(row=k, column=col).font = base_font
        return start_row + len(values) + 2

    pairs = sorted({r["pair"] for r in rows})
    tfs = sorted({r["timeframe"] for r in rows if r["timeframe"]})
    nxt = breakdown(19, "By pair", pair_rng, pairs or ["—"])
    if tfs:
        breakdown(nxt, "By timeframe", tf_rng, tfs)

    # ---- Equity sheet
    we = wb.create_sheet("Equity")
    money_fmt = "#,##0.00;[Red]-#,##0.00"
    for j, h in enumerate(["Closed", "Pair", "P&L ($)", "Cumulative ($)", "Peak ($)", "Drawdown ($)"], start=1):
        c = we.cell(row=1, column=j, value=h)
        c.font, c.fill = head_font, head_fill
    for col, w in zip("ABCDEF", (18, 10, 11, 14, 11, 13)):
        we.column_dimensions[col].width = w
    closed = sorted([r for r in rows if r["result"] in CLOSED and r["pnl"] is not None],
                    key=lambda x: x["closed"] or x["opened"] or datetime.min)
    for k, r in enumerate(closed, start=2):
        we.cell(row=k, column=1, value=r["closed"] or r["opened"]).number_format = "yyyy-mm-dd hh:mm"
        we.cell(row=k, column=2, value=r["pair"])
        we.cell(row=k, column=3, value=r["pnl"]).number_format = money_fmt
        we.cell(row=k, column=4, value=f"=C{k}" if k == 2 else f"=D{k-1}+C{k}").number_format = money_fmt
        # running peak starts from 0 (the starting balance)
        we.cell(row=k, column=5, value=f"=MAX(0,D{k})" if k == 2 else f"=MAX(E{k-1},D{k})").number_format = money_fmt
        we.cell(row=k, column=6, value=f"=E{k}-D{k}").number_format = "#,##0.00"
    end_row = len(closed) + 1
    if closed:
        chart = LineChart()
        chart.title = "Equity curve (cumulative P&L)"
        chart.y_axis.title = "USD"
        chart.x_axis.title = "Trade #"
        chart.height, chart.width = 9, 22
        chart.legend = None
        chart.add_data(Reference(we, min_col=4, min_row=1, max_row=end_row), titles_from_data=True)
        we.add_chart(chart, "H2")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------- PDF
def build_pdf(rows: list, period: str) -> bytes:
    from reportlab.graphics.charts.lineplots import LinePlot
    from reportlab.graphics.shapes import Drawing, Line, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)

    s = summarize(rows)
    styles = getSampleStyleSheet()
    navy = colors.HexColor("#1F2A44")
    green, red = colors.HexColor("#15803D"), colors.HexColor("#B91C1C")
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm,
                            title="Forex Intel — Trade Journal")
    story = [
        Paragraph("Forex Intel — Trade Journal Report", styles["Title"]),
        Paragraph(f"Period: {period} &nbsp;·&nbsp; Exported {datetime.now():%Y-%m-%d %H:%M}", styles["Normal"]),
        Spacer(1, 6 * mm),
    ]

    def money(v):
        if v is None:
            return "—"
        return f"-${abs(v):,.2f}" if v < 0 else f"${v:,.2f}"

    def price(pair, v):
        if v is None:
            return ""
        dp = 3 if "JPY" in pair else 2 if "XAU" in pair else 5
        return f"{v:.{dp}f}"

    kpis = [
        ["Trades", "Closed", "Wins / Losses", "Win rate", "Net P&L", "Profit factor", "Avg R", "Max drawdown", "Worst losing run"],
        [s["total"], s["closed"], f"{s['wins']} / {s['losses']}", f"{s['win_rate']:.1f}%", money(s['net']),
         "—" if s["profit_factor"] is None else f"{s['profit_factor']:.2f}",
         "—" if s["avg_r"] is None else f"{s['avg_r']:+.2f}", money(s['max_dd']), s["worst_streak"]],
    ]
    t = Table(kpis, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"), ("FONTSIZE", (0, 1), (-1, 1), 12),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("GRID", (0, 0), (-1, -1), 0.4, colors.lightgrey),
        ("TOPPADDING", (0, 1), (-1, 1), 6), ("BOTTOMPADDING", (0, 1), (-1, 1), 6),
        ("TEXTCOLOR", (4, 1), (4, 1), green if s["net"] >= 0 else red),
    ]))
    story += [t, Spacer(1, 2 * mm),
              Paragraph("<font size=8 color='#666666'>Break-even win rate at 2:1 is about 33% "
                        "(a bit higher after costs). R = result divided by the money risked.</font>",
                        styles["Normal"]),
              Spacer(1, 5 * mm)]

    # per pair
    closed = [r for r in rows if r["result"] in CLOSED and r["pnl"] is not None]
    by_pair = [["Pair", "Trades", "Wins", "Losses", "Win rate", "Net P&L ($)", "Avg R"]]
    for p in sorted({r["pair"] for r in closed}):
        pr = [r for r in closed if r["pair"] == p]
        w = sum(1 for r in pr if r["result"] == "WIN")
        rs = [r["r"] for r in pr if r["r"] is not None]
        by_pair.append([p, len(pr), w, sum(1 for r in pr if r["result"] == "LOSS"),
                        f"{w / len(pr) * 100:.0f}%", f"{sum(r['pnl'] for r in pr):+,.2f}",
                        f"{sum(rs) / len(rs):+.2f}" if rs else "—"])
    left = Table(by_pair, hAlign="LEFT")
    left_style = [("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                  ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 9),
                  ("GRID", (0, 0), (-1, -1), 0.4, colors.lightgrey), ("ALIGN", (1, 0), (-1, -1), "RIGHT")]
    for i, row in enumerate(by_pair[1:], start=1):
        left_style.append(("TEXTCOLOR", (5, i), (5, i), green if not row[5].startswith("-") else red))
    left.setStyle(TableStyle(left_style))

    # equity curve
    pts, eq = [(0, 0.0)], 0.0
    for i, r in enumerate(sorted(closed, key=lambda x: x["closed"] or x["opened"] or datetime.min), start=1):
        eq += r["pnl"]
        pts.append((i, round(eq, 2)))
    d = Drawing(130 * mm, 62 * mm)
    d.add(String(0, 58 * mm, "Equity curve (cumulative P&L, USD)", fontName="Helvetica-Bold", fontSize=9))
    if len(pts) > 1:
        lp = LinePlot()
        lp.x, lp.y, lp.width, lp.height = 12 * mm, 8 * mm, 112 * mm, 46 * mm
        lp.data = [pts]
        lp.lines[0].strokeColor = colors.HexColor("#4F46E5")
        lp.lines[0].strokeWidth = 1.5
        ys = [p[1] for p in pts]
        lo, hi = min(ys), max(ys)
        pad = max((hi - lo) * 0.1, 1)
        lp.yValueAxis.valueMin, lp.yValueAxis.valueMax = lo - pad, hi + pad
        lp.xValueAxis.valueMin, lp.xValueAxis.valueMax = 0, max(len(pts) - 1, 1)
        lp.xValueAxis.labels.fontSize = lp.yValueAxis.labels.fontSize = 7
        d.add(lp)
        if lo < 0 < hi:
            y0 = lp.y + (0 - (lo - pad)) / ((hi + pad) - (lo - pad)) * lp.height
            d.add(Line(lp.x, y0, lp.x + lp.width, y0, strokeColor=colors.grey, strokeDashArray=[2, 2]))
    else:
        d.add(String(0, 30 * mm, "No closed trades in this period.", fontSize=9))
    side = Table([[left, d]], colWidths=[125 * mm, 135 * mm], hAlign="LEFT")
    side.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [side, PageBreak()]

    # all trades
    story.append(Paragraph("All trades", styles["Heading2"]))
    head = ["Opened", "Pair", "Dir", "TF", "Entry", "SL", "TP", "Close", "Lots", "Risk $", "Result", "P&L $", "R", "Closed by"]
    data = [head]
    for r in rows:
        data.append([
            f"{r['opened']:%d %b %H:%M}" if r["opened"] else "", r["pair"], r["direction"], r["timeframe"],
            price(r["pair"], r["entry"]), price(r["pair"], r["stop_loss"]),
            price(r["pair"], r["take_profit"]), price(r["pair"], r["close"]),
            r["lots"], f"{r['risk']:.2f}", r["result"],
            "" if r["pnl"] is None else f"{r['pnl']:+.2f}",
            "" if r["r"] is None else f"{r['r']:+.2f}", r["close_reason"],
        ])
    tt = Table(data, repeatRows=1, hAlign="LEFT")
    ts = [("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
          ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7.5),
          ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
          ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6FA")])]
    for i, r in enumerate(rows, start=1):
        if r["result"] == "WIN":
            ts.append(("TEXTCOLOR", (10, i), (12, i), green))
        elif r["result"] == "LOSS":
            ts.append(("TEXTCOLOR", (10, i), (12, i), red))
    tt.setStyle(TableStyle(ts))
    story.append(tt)

    def footer(canvas, _doc):
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(14 * mm, 7 * mm, "Forex Intel — decision-support tool, not financial advice.")
        canvas.drawRightString(landscape(A4)[0] - 14 * mm, 7 * mm, f"Page {_doc.page}")

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ------------------------------------------------------------------- routes
@router.get("/excel")
def export_excel(
    start: Optional[date] = Query(default=None),
    end: Optional[date] = Query(default=None),
    pair: Optional[str] = Query(default=None),
    result: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    rows = load_rows(db, current_user, start, end, pair, result)
    data = build_excel(rows, _period(rows, start, end))
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_filename("xlsx")}"'},
    )


@router.get("/pdf")
def export_pdf(
    start: Optional[date] = Query(default=None),
    end: Optional[date] = Query(default=None),
    pair: Optional[str] = Query(default=None),
    result: Optional[str] = Query(default=None),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    rows = load_rows(db, current_user, start, end, pair, result)
    data = build_pdf(rows, _period(rows, start, end))
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{_filename("pdf")}"'},
    )
